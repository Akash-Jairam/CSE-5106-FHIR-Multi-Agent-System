"""Build the offline V3 handoff from saved numerical records; never calls a model."""
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1]
REPO = BASE.parents[1]
load = lambda p: json.loads(p.read_text())
s = load(BASE / 'full_single_agent/single_agent_summary.json')
v = load(BASE / 'full_single_agent/matched_v3_summary.json')
pair = load(BASE / 'full_single_agent/paired_comparison.json')
dev = load(BASE / 'development/statistics.json')
families = load(BASE / 'full_single_agent/by_family.json')
paired_families = load(BASE / 'full_single_agent/paired_by_family.json')
overlap = load(BASE / 'full_single_agent/exploratory_development_exposure_overlap.json')
assert s['questions'] == v['questions'] == pair['matched_questions'] == 2929
assert (s['correct'], v['correct']) == (2343, 2255)
assert overlap['known_union_patients'] == 93

def metric(arm, key):
    return arm['metrics'][key]

def pct(x):
    return f'{100*x:.2f}%'

def usd(x, n=6):
    return f'${x:,.{n}f}'

def ci(o, scale=1, digits=2):
    return f"[{o['ci_low']*scale:,.{digits}f}, {o['ci_high']*scale:,.{digits}f}]"

cost_sensitivity = {'cohort_n': 2929, 'interpretation': 'Fixed-usage accounting scenario, not another inference experiment.', 'arms': {}}
for name, arm in [('single_agent', s), ('multi_agent_v3', v)]:
    m = arm['metrics']
    p, k, c = [m[key]['total'] for key in ['prompt_tokens', 'cached_prompt_tokens', 'completion_tokens']]
    nc = (p*.10+c*.50)/1e6
    cost_sensitivity['arms'][name] = {
        'prompt_tokens': p, 'cached_prompt_tokens': k, 'completion_tokens': c,
        'cached_prompt_fraction': k/p, 'recorded_cache_total_usd': m['estimated_cost_usd']['total'],
        'no_cache_total_usd': nc, 'no_cache_mean_usd': nc/2929,
    }
cost_sensitivity['v3_to_single_recorded_cache_cost_ratio'] = metric(v, 'estimated_cost_usd')['total']/metric(s, 'estimated_cost_usd')['total']
cost_sensitivity['v3_to_single_no_cache_cost_ratio'] = cost_sensitivity['arms']['multi_agent_v3']['no_cache_total_usd']/cost_sensitivity['arms']['single_agent']['no_cache_total_usd']
(BASE / 'cost_cache_sensitivity.json').write_text(json.dumps(cost_sensitivity, indent=2)+'\n')

registry = {
 'schema_version': 1, 'branch': 'results/v3-ful-run', 'date_utc': '2026-10-17',
 'analysis_revised_date_utc': '2026-10-19',
 'purpose': 'Four historical arm records and a primary full benchmark comparison of 2,929 completed task pairs; no new inference.',
 'runs': [
  {'id':'development_single_agent_300','arm':'single_agent','accepted_or_scheduled_n':300,'planned_n':300,'correct':238,'policy':'historical_first_attempt','results':'results/v3/development','original_results':'benchmarks/design-v3d-20261014','record_commit':'906327dee1823178e3ac6319061f7d9ee5d2e01f','source_commit':'f00804f186768a3672ae6a6075f366a27dc0ca9a'},
  {'id':'development_v3_300','arm':'multi_agent_v3','accepted_or_scheduled_n':300,'planned_n':300,'correct':232,'policy':'historical_first_attempt','results':'results/v3/development','original_results':'benchmarks/design-v3d-20261014','record_commit':'906327dee1823178e3ac6319061f7d9ee5d2e01f','source_commit':'f00804f186768a3672ae6a6075f366a27dc0ca9a'},
  {'id':'full_v3_2931','arm':'multi_agent_v3','accepted_or_scheduled_n':2931,'planned_n':2931,'correct':2256,'policy':'first_clean_attempt','complete':True,'results':'results/v3/full_v3','original_results':'benchmarks/v3-full-run-20261015/completed_v3','source_commit':'4c7fbd13c1fd6e0b38abd002b072af8654837f28'},
  {'id':'full_single_agent_available_2929','arm':'single_agent','accepted_or_scheduled_n':2929,'planned_n':2931,'correct':2343,'policy':'first_clean_attempt','complete':False,'results':'results/v3/full_single_agent','source_commit':'4c7fbd13c1fd6e0b38abd002b072af8654837f28'},
 ],
 'primary_comparison': {'type':'posthoc_available_matched_cohort','report_label':'Full benchmark comparison of 2,929 completed task pairs','n':2929,'patients':94,'single_correct':2343,'v3_correct':2255,'results':'results/v3/full_single_agent/paired_comparison.json','summary':'results/v3/comparison_summary.csv','difference_convention':'full comparison paired files Single Agent minus V3; development paired tables and handoff V3 minus Single Agent'},
 'overlap': {'development_question_ids_in_matched_full':300,'fresh_development_patients':58,'known_union_exposed_patients':93,'full_patients':94,'independent_repetitions':False,'untouched_holdout_claim':False},
 'source_comparison':'results/v3/development/source_comparison.json',
 'source_note':'Full and development share V3 architecture; host-loop bound21 versus22 differs in one shared runtime file. Model-call cap20 unchanged.',
 'runtime_implementation_merged':False,
}
(BASE / 'RUN_REGISTRY.json').write_text(json.dumps(registry, indent=2)+'\n')

doc = Document()
# The bundled base template can carry a title rule; keep headings plain.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.7), Inches(.7)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)
styles = doc.styles
for name in ['Normal','Body Text']:
    st=styles[name]; st.font.name='Calibri'; st.font.size=Pt(11)
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.08
for name, size in [('Title',25),('Heading 1',18),('Heading 2',12)]:
    st=styles[name]; st.font.name='Calibri'; st.font.size=Pt(size)
    st.font.color.rgb=RGBColor(0,0,0); st.font.bold=True
    st.paragraph_format.space_before=Pt(10)
    st.paragraph_format.space_after=Pt(9)
    st.paragraph_format.keep_with_next=True
header = sec.header.paragraphs[0]
header.text='V3 MIDTERM RESULTS   |   2,929 completed pairs   |   Revised 19 October 2026'
header.runs[0].font.size=Pt(8)
header.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Results branch  results/v3-ful-run   |   '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
text=[]

def p(value, bold=False):
    par=doc.add_paragraph(); run=par.add_run(value); run.bold=bold
    text.append(value+'\n')
    return par

def h(value, level=1):
    doc.add_heading(value, level=level); text.append('\n'+value.upper()+'\n')

def page(value):
    doc.add_page_break(); h(value)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths is None: widths=[7.1/len(headers)]*len(headers)
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for i,value in enumerate(headers): t.rows[0].cells[i].text=str(value)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row): cells[i].text=str(value)
    for row_no,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); pr.append(el)
        if row_no==0:
            repeat=OxmlElement('w:tblHeader'); pr.append(repeat)
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'17365D' if row_no==0 else ('F1F5F9' if row_no%2 else 'FFFFFF')); tcpr.append(shd)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge,val in [('top',75),('bottom',75),('left',85),('right',85)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(1);par.paragraph_format.line_spacing=1.02
                par.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 or headers[-1]=='Use' else WD_ALIGN_PARAGRAPH.RIGHT
                for run in par.runs:
                    run.font.size=Pt(9.5);run.font.bold=(row_no==0)
                    run.font.color.rgb=RGBColor(255,255,255) if row_no==0 else RGBColor(0,0,0)
    text.append(' | '.join(headers)+'\n'+'\n'.join(' | '.join(map(str,r)) for r in rows)+'\n')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

doc.add_paragraph('V3 consolidated results',style='Title')
p('Midterm checkpoint and analysis handoff',True)
p('Development checkpoint and full benchmark comparison • Revised 19 October 2026')
p('We ran both systems against the full benchmark and report comparative results for 2,929 completed task pairs from 94 patients. Single Agent has higher answer accuracy, lower task latency and lower estimated cost. V3 uses fewer measured tokens and tool calls. These results describe the tradeoff in the current V3 design.')
table(['Analysis record','Questions','Correct','Accuracy'],[
 ['Development Single Agent','300','238','79.33%'],['Development V3','300','232','77.33%'],
 ['Full comparison Single Agent','2,929','2,343','79.99%'],['Full comparison V3','2,929','2,255','76.99%']], [3.15,1.2,1.2,1.55])
h('Main comparison',2)
p('Use the same 2,929 completed question IDs for both arms. On this cohort V3 scores 2,255/2,929 (76.99%). V3 minus Single Agent is −3.00 percentage points, with a paired patient-cluster 95% interval of [−4.17, −1.84] percentage points.')
p('V3 mean task latency is 24.06 seconds versus 17.93 seconds; estimated mean cost is $0.005055 versus $0.002314. V3 uses 53,512 versus 73,273 tokens per question, about 27.0% fewer. Recorded cached-input shares differ substantially, so fewer tokens do not imply a lower dollar estimate.')
h('How to use this handoff',2)
p('All full benchmark accuracy, time, token and cost figures in this document use the same 2,929 completed question IDs for both systems. Development uses the historical 300-question first-attempt protocol and overlaps the full benchmark. The four analysis rows describe two systems at two stages; they are not independent repetitions. Original execution counts remain in the archive registry.')
p('The remainder explains methods, development findings, paired accuracy, time and usage, cost sensitivity, mechanisms, limitations and file-level reproduction. Detailed machine-readable statistics accompany this document.')

page('Design and comparison methods')
h('Architecture under evaluation',2)
p('V3 is the question-contract architecture, recorded as query-contract-v3.2-static-audit. A question-only contract stage has no access to the solver answer, source records or gold answer. A continuous solver retrieves and computes against FHIR data with an execution ledger and source validation. A static program audit checks the contract, executed code and bounded sources without receiving a rendered candidate answer. Supported host repair is narrowly constrained; the design does not run a second full solve or vote.')
p('Recorded phase ceilings are 10,000 contract tokens, 175,000 solver tokens and 15,000 audit tokens, within an aggregate 200,000-token and 30-executed-tool budget. The solver has at most 20 model calls. The Single Agent is the contemporaneous implementation control; an earlier paper score would not measure this model, substrate, scorer, cache behavior and question cohort together.')
h('Frozen execution and scoring',2)
p('The recorded Azure deployment is gpt-6-luna, reasoning none, with a maximum completion limit of 16,384 tokens, a 32,000-token estimated-input guard and enabled prompt caching. Full runs use two workers. The unchanged typed-answer scorer determines correctness from saved final answers and reference types. No model judge or additional inference is used for consolidation.')
p('Full runs select the first accepted runtime-clean task outcome by the frozen collection policy. Only those selected task measurements enter the design totals. Development retains its original scheduled first-attempt scores. Unknown measurement values remain null. Missing fields are not converted to free usage or zero time.')
h('Cohort and uncertainty',2)
p('The primary comparison joins completed question IDs and verifies patient IDs, giving 2,929 matched pairs and 94 patients. Pairing was determined by outcome availability, without selecting by correctness; this is a posthoc descriptive comparison. Patient-cluster percentile intervals resample patients 10,000 times with seed 5106, retain their questions and estimate the question-weighted mean. Pairing is preserved in differences. Intervals describe these questions and the recorded run, not repeated deployments or random model seeds.')
p('The dataset informed development. All 300 fresh development question IDs also occur in the matched full cohort. The original split labels do not capture later development exposure; they must not be called an untouched holdout. See the exposure discussion on page 8.')
h('Source comparability',2)
p('The full source freeze is 4c7fbd13; the development source freeze is f00804f1, with results pinned at 906327de. Of 57 declared shared runtime files, 56 are identical. LedgerSolver has a host-loop bound of 21 in the full freeze and 22 in development; its actual model-call cap stays 20. This can affect final-boundary completion, so the runs share the V3 architecture but are not byte-identical. Consolidation merges results, not development runtime code.')

page('Development checkpoint')
p('The fresh paired development cohort contains 300 questions from 58 patients. Single Agent scores 238/300 and V3 scores 232/300. V3 minus Single Agent is −2.00 percentage points; the frozen paired 95% interval [−5.16, +1.10] includes zero. This reused development sample does not establish an accuracy advantage for V3.')
table(['Question family','N','Single correct','V3 correct'],[
 ['Lookup and counting','92','81 (88.04%)','79 (85.87%)'],['Numerical aggregation','49','40 (81.63%)','37 (75.51%)'],['Temporal selection','159','117 (73.58%)','116 (72.96%)'],['All','300','238 (79.33%)','232 (77.33%)']], [2.65,.65,1.9,1.9])
p('The paired outcomes are 224 both correct, 54 both wrong, 8 V3-only correct and 14 Single-only correct. The fresh contemporaneous control is the relevant comparison here; older V1 or refreshed V2 control scores must not be substituted.')
table(['Measurement','Single Agent','V3'],[
 ['Task latency total','5,413.751 s','5,297.728 s'],['Task latency mean','18.046 s','17.659 s'],['Task latency median','15.875 s','15.140 s'],['Task latency p95','34.640 s','35.588 s'],['Measured tokens total','21,057,621','15,507,699'],['Measured tokens mean','70,192','51,692'],['Model calls mean','6.487','5.993'],['Executed tools mean','6.503','5.150']], [3.1,2,2])
p('The paired mean latency difference is −0.387 seconds, 95% interval [−1.463, +0.699]. Measured token reduction is 18,499.74 per question, interval [14,706.89, 22,184.19] fewer tokens. V3 therefore reduced measured resource usage while showing no clear development latency or accuracy gain.')
h('Development cost and behavior',2)
p('Single Agent has complete recorded question-level usage for 300/300, giving a $0.671200 total proxy cost and $0.002237 mean. V3 has complete whole-question usage for 298/300: the observed sum is $0.873795 and observed mean $0.002932. Its full 300-question cost total is unavailable. On the same 298 observed pairs, the V3-minus-Single mean cost difference is +$0.000690, interval [+$0.000419, +$0.000953].')
p('V3 produced 283 source-valid primary releases. Source feedback triggered 19 revisions with four wrong-to-right transitions and no losses at that specific transition. The static audit checked 283 valid primaries, proposed no material local repair and triggered zero targeted revisions. These diagnostics motivate examining the audit contribution; they do not identify a causal effect.')

page('Full paired answer accuracy')
p('The primary full comparison uses 2,929 matched questions from 94 patients. Single Agent is correct on 2,343 (79.99%, individual-arm 95% interval [77.20%, 82.49%]); V3 is correct on 2,255 (76.99%, interval [74.54%, 79.37%]). The paired V3-minus-Single difference is −3.00 percentage points, interval [−4.17, −1.84].')
table(['Outcome','Questions','Share'],[
 ['Both correct','2,162',pct(2162/2929)],['Single Agent only correct','181',pct(181/2929)],['V3 only correct','93',pct(93/2929)],['Both wrong','493',pct(493/2929)],['All matched','2,929','100.00%']], [4.1,1.5,1.5])
p('The net difference is 88 questions: V3 recovers 93 questions missed by Single Agent but misses 181 that Single Agent answers correctly. A future design should investigate both directions, rather than examining only examples where cooperation succeeds. The saved paired question IDs allow those cases to be recovered from the raw envelopes without mixing in the development record.')
h('Question family comparison',2)
family_rows=[]
labels={'lookup_counting':'Lookup and counting','numerical_aggregation':'Numerical aggregation','temporal_selection':'Temporal selection','out_of_scope':'Outside core families'}
for f in paired_families:
    group=f['group']; n=f['matched_questions']; sc=f['both_correct']+f['left_wins'];vc=f['both_correct']+f['right_wins']
    d=f['accuracy_difference'];
    family_rows.append([labels.get(group,group),str(n),f'{100*sc/n:.2f}%',f'{100*vc/n:.2f}%',f"{-100*d['estimate']:+.2f}"])
table(['Family','N','Single','V3','V3 delta pp'],family_rows,[2.8,.65,1.1,1.1,1.45])
p('The largest point deficit is temporal selection: −4.42 percentage points, interval [−6.25, −2.70]. Lookup and counting is −3.45 points, interval [−5.50, −1.46]. Numerical aggregation is −2.96 points, interval [−6.91, +0.92]; the outside-core group is +1.74 points, interval [−1.76, +5.32]. These subgroup intervals are exploratory and are not multiplicity-adjusted.')
p('Family labels are the frozen analysis taxonomy. Outside-core questions remain in the overall denominator because they belong to the released benchmark. Removing them or selecting only answered questions would change the target quantity and should be declared as a separate analysis.')
h('Answer coverage is a separate diagnostic',2)
p('Saved empty/no-value flags can occur when an empty answer is correct. They are not a refusal classifier. Full matched records contain 548 flagged Single Agent outputs and 653 flagged V3 outputs; the accompanying correctness cross-tabulation preserves which are correct. Use overall typed accuracy as the principal answer metric.')

page('Task time and token usage')
rows=[]
for key,label,fmt in [
 ('latency_seconds','Task latency seconds',lambda x:f'{x:,.3f}'),
 ('total_tokens','Total tokens',lambda x:f'{x:,.0f}'),
 ('llm_calls','Model calls',lambda x:f'{x:,.3f}'),
 ('tool_calls_executed','Executed tools',lambda x:f'{x:,.3f}')]:
    for stat,title in [('mean','mean'),('median','median'),('p95','p95')]:
        rows.append([f'{label} {title}',fmt(metric(s,key)[stat]),fmt(metric(v,key)[stat])])
table(['Matched cohort measurement','Single Agent','V3'],rows,[3.7,1.7,1.7])
p('V3 minus Single Agent mean latency is +6.127 seconds, 95% paired interval [+5.303, +6.943]. The ratio of means is 1.342: V3 takes about 34.2% longer per selected task. Median and p95 latencies also favor Single Agent. The distributions are skewed; both median and tail measurements should accompany means in a presentation.')
p('Accepted task latencies sum to 52,531.060 seconds for Single Agent and 70,478.026 seconds for V3 on the matched cohort. These are accumulated task times, not elapsed experiment durations. The two-worker schedule and separate execution dates prevent converting their difference into a controlled wall-clock speedup.')
p('V3 consumes 156,736,271 total tokens versus 214,615,829, a reduction of 57,879,558 tokens or 27.0%. The paired per-question reduction is 19,760.86 tokens, interval [16,361.59, 23,490.72]. Model-call totals are 17,708 versus 19,440; executed-tool totals are 15,502 versus 19,473. Fewer calls and tools coexist with greater task latency, so call counts alone do not explain elapsed time.')
h('Spread and completeness',2)
p('Latency sample standard deviations are 9.965 seconds for Single Agent and 13.982 seconds for V3. Observed ranges are [2.876, 114.961] and [5.692, 141.772] seconds. Total-token sample standard deviations are 47,527.98 and 37,883.92 respectively. All matched primary latency, token and estimated-cost values are present.')
p('The detailed tables retain prompt, completion, cached and uncached prompt tokens, call-ordinal usage, phases where recorded and tool-name counts. The repeated-prompt field is a sequence-based heuristic, not exact text overlap or an added billing category. Missing phase labels do not mean that a phase consumed zero resources.')

page('Cost and cache accounting')
p('All dollar amounts are estimates under the user-approved frozen official OpenAI proxy for Azure: $0.10 per million uncached input tokens, $0.01 per million cached input tokens and $0.50 per million output tokens. Actual invoices are unavailable. Cached tokens are included in prompt totals and are not charged twice.')
p('Estimated cost = ((prompt − cached) × 0.10 + cached × 0.01 + completion × 0.50) / 1,000,000.',True)
table(['Matched cohort cost','Single Agent','V3'],[
 ['Total estimated USD',usd(metric(s,'estimated_cost_usd')['total']),usd(metric(v,'estimated_cost_usd')['total'])],
 ['Mean per question',usd(metric(s,'estimated_cost_usd')['mean']),usd(metric(v,'estimated_cost_usd')['mean'])],
 ['Median per question',usd(metric(s,'estimated_cost_usd')['median']),usd(metric(v,'estimated_cost_usd')['median'])],
 ['p95 per question',usd(metric(s,'estimated_cost_usd')['p95']),usd(metric(v,'estimated_cost_usd')['p95'])],
 ['Sample standard deviation',usd(metric(s,'estimated_cost_usd')['stdev']),usd(metric(v,'estimated_cost_usd')['stdev'])],
 ['Cached share of prompt tokens',pct(cost_sensitivity['arms']['single_agent']['cached_prompt_fraction']),pct(cost_sensitivity['arms']['multi_agent_v3']['cached_prompt_fraction'])],
 ['Uncached prompt tokens','29,496,435','125,572,340']], [3.5,1.8,1.8])
p('The paired V3-minus-Single mean cost difference is +$0.002741, interval [+$0.002573, +$0.002906]. The observed ratio is 2.185, or 118.5% greater estimated cost for V3. Most of the contrast comes from the cache mix: V3 uses fewer prompt tokens overall but about 4.26 times as many uncached prompt tokens.')
h('Fixed usage sensitivity',2)
cs=cost_sensitivity['arms']['single_agent'];cv=cost_sensitivity['arms']['multi_agent_v3']
table(['Accounting scenario','Single total USD','V3 total USD'],[
 ['Recorded cache mix',usd(cs['recorded_cache_total_usd']),usd(cv['recorded_cache_total_usd'])],
 ['All prompt tokens at uncached price',usd(cs['no_cache_total_usd']),usd(cv['no_cache_total_usd'])]], [3.5,1.8,1.8])
p(f"If the same measured usage were priced with no cached-input discount, V3's estimated total would be {100*(1-cost_sensitivity['v3_to_single_no_cache_cost_ratio']):.1f}% lower. This scenario isolates the accounting contribution of caching; it is not another inference run and does not show what latency or outputs would have been without caching.")
p('Cache availability flags require care. The V3 cache reconciliation covers every V3 question used in this comparison. The preserved baseline provenance checkpoint covers 891 questions; a full baseline cache-metadata audit was not performed. Normalized numeric cache values are sufficient for reproducing these estimates, but are not a guarantee that every original metadata field was present. Dates, scheduling and prompt-prefix reuse were not randomized between arms.')

page('Mechanism diagnostics')
p('The saved diagnostics distinguish final answer correctness from source support, draft quality and local revision behavior. They are useful for explaining where the design may be losing value, while observational transitions cannot isolate the effect of a particular component.')
table(['Matched full diagnostic','Single Agent','V3'],[
 ['Questions with a saved draft','2,806','2,912'],['Correct saved drafts','2,220','2,135'],
 ['Draft wrong to final correct','96','171'],['Draft correct to final wrong','15','56'],
 ['Final correct among draft-present','2,301','2,250'],['Final correct among draft-absent','42','5'],
 ['All final correct','2,343','2,255'],['Source-check records available','2,806','2,929'],
 ['Source checks marked clean','2,668','2,640'],['Solver source revisions','0','351']], [3.5,1.8,1.8])
p('Draft arithmetic reconciles within its actual scope. For Single Agent, 2,220 + 96 − 15 = 2,301 final-correct draft-present questions; 42 additional correct outputs have no saved draft. V3 has 2,135 + 171 − 56 = 2,250, plus five correct outputs without drafts. This prevents treating absent drafts as incorrect drafts or comparing stage totals to incompatible final denominators.')
p('V3 improves more wrong drafts in absolute count, but it starts with fewer correct saved drafts and also loses more correct drafts. Net draft-present gains are +81 for Single Agent and +115 for V3. These counts are compatible with V3 having lower final accuracy and cannot be summarized simply as a successful revision rate.')
h('Source and audit behavior',2)
p('V3 records 2,640 releases under source_valid_primary_preserved and 289 under primary_failed_execution_or_source_contract. Its 351 solver source revisions and 299 answer changes describe different events and need not be equal. The matched full mechanism export records no targeted revisions or targeted wrong-primary recoveries. Together with development’s zero material audit repairs, this provides little observed evidence that the static audit supplies the intended answer-level gain.')
p('Matched V3 evidence counts total 19,297, with 19,252 marked valid by recorded local checks. Evidence validation is unavailable for Single Agent and must not be filled with zero or compared as an accuracy score. Local source validity does not establish temporal interpretation, clinical meaning or final-answer correctness.')
p('The baseline records 176 budget-stop flags and V3 zero. A flag is an implementation diagnostic, not a substitute for scoring the saved answer. Mechanism fields, source checks, limitations and evidence precision/recall have different missingness and denominators; their full tables preserve that information.')

page('Exposure and interpretation limits')
h('Development exposure is broader than source split labels',2)
p('The older split labels identify 35 dev/pilot patients and 59 patients labeled held_out_patient_descriptive. The later fresh 300-question development run uses 58 patients, all drawn from that latter label. Those 58 patients contribute 1,678 matched full questions; all 300 exact development question IDs also occur in the matched cohort. The known union of the two development sources covers 93 of 94 full patients.')
p('Consequently, the saved exposure subgroup tables are source-split descriptions, not evidence of untouched generalization. Their frozen labels are retained for reproducibility, and a separate exploratory overlap file records the later exposure. This handoff does not make a heldout-performance claim or construct a new confirmatory subgroup from the remaining patient.')
h('What the experiment establishes',2)
p('For this frozen V3 architecture and these recorded runs on the released benchmark, V3 reduces measured token and tool usage. Across the 2,929 completed task pairs it has lower typed accuracy, greater per-task latency and greater cost under the recorded cache mix than the contemporaneous Single Agent control. The accuracy interval is conditional on patient sampling from this benchmark; it does not account for design selection on reused data.')
p('A controlled causal claim about multi-agent architecture alone would require tighter isolation. Execution dates and cache states differ, the development and full code have a documented host-loop bound difference, and the cohort informed refinement. The no-cache price calculation demonstrates accounting sensitivity without simulating a new service configuration.')
h('What should travel with a midterm figure',2)
p('Any accuracy figure should state the question denominator, 94-patient clustering for the full comparison, reused development data and paired interval method. Any efficiency figure should name its metric: selected task seconds, measured tokens or estimated USD. Cost figures should state the rates and the cache assumption. Keep the 300-question checkpoint separate from the full common-cohort comparison.')
p('Do not replace typed accuracy with conditional accuracy among nonempty outputs; do not add token counts to proxy dollars; do not treat task-time sums as wall-clock run duration; and do not infer independent evidence validation where it was not recorded. These distinctions affect interpretation directly.')
h('Questions supported for later work',2)
p('The paired discordant question IDs support analyzing the 181 Single-only and 93 V3-only successes by temporal interpretation, source-contract enforcement and computation behavior. The draft transitions support studying why correct drafts are lost and why audit repair contributes no observed targeted recovery. These are follow-up analyses or new protocol decisions; this consolidation launches no additional benchmark.')

page('V3 comparison reference')
p(f"V3 is evaluated on the same {v['questions']:,} completed question IDs as Single Agent, from 94 patients: {v['correct']:,} correct and {v['questions']-v['correct']:,} wrong, or {100*v['accuracy']['estimate']:.4f}% accuracy. Its patient-cluster 95% interval is {ci(v['accuracy'],100,4)}%. All values below use these 2,929 questions and agree with the paired analysis.")
vm=v['metrics']
table(['V3 metric for 2929 questions','Value'],[
 ['Accuracy',f"{v['correct']:,} / {v['questions']:,} = {100*v['accuracy']['estimate']:.4f}%"],
 ['Task latency sum',f"{vm['latency_seconds']['total']:,.3f} s"],
 ['Task latency mean / median',f"{vm['latency_seconds']['mean']:.4f} s / {vm['latency_seconds']['median']:.4f} s"],
 ['Task latency p95 / sample SD',f"{vm['latency_seconds']['p95']:.4f} s / {vm['latency_seconds']['stdev']:.4f} s"],
 ['Mean latency 95% interval',ci(vm['latency_seconds']['mean_patient_bootstrap'],digits=4)+' s'],
 ['Estimated total cost',usd(vm['estimated_cost_usd']['total'],8)],
 ['Estimated mean / median cost',usd(vm['estimated_cost_usd']['mean'],8)+' / '+usd(vm['estimated_cost_usd']['median'],8)],
 ['Estimated p95 cost',usd(vm['estimated_cost_usd']['p95'],9)],
 ['Mean cost 95% interval',ci(vm['estimated_cost_usd']['mean_patient_bootstrap'],digits=8)+' USD'],
 ['Total measured tokens',f"{vm['total_tokens']['total']:,}"],
 ['Mean / median total tokens',f"{vm['total_tokens']['mean']:,.2f} / {vm['total_tokens']['median']:,.0f}"],
 ['p95 total tokens',f"{vm['total_tokens']['p95']:,.1f}"],
 ['Mean tokens 95% interval',ci(vm['total_tokens']['mean_patient_bootstrap'])],
 ['Primary measurement missingness','0 latency / cost estimate / tokens']], [3.65,3.45])
p('The source summaries are full_single_agent/matched_v3_summary.json and single_agent_summary.json. The root comparison_summary.csv provides both systems on the same denominator for later report tables and figures. Preserve these question IDs when joining further diagnostics.')
h('Correct empty answers remain correct',2)
p('Among the 2,929 V3 comparison outputs, 653 have the syntactic empty/no-value flag. Of these, 422 are correct against empty reference results and 231 are incorrect. Among the other 2,276 outputs, 1,833 are correct and 443 incorrect. The 76.9887% score includes legitimate empty answers. This distinction is preserved in the numeric records.')
p('Use the 70,478.026-second selected task-time sum for this V3 comparison. Historical experiment wall intervals remain in the execution archive; concurrency and scheduling make them different measurements from accumulated task time.')

page('Files metrics and reproduction')
p('All paths below are relative to the repository root. The branch results/v3-ful-run contains the historical result directories, derived numerical tables, this editable document, a text counterpart and portable raw archives. No development runtime implementation is merged.')
table(['Location under results/v3','Use'],[
 ['RUN_REGISTRY.json and METHODS.txt','Run identities, cohorts and analysis protocol'],
 ['comparison_summary.csv','Primary V3 and Single Agent figures on 2,929 questions each'],
 ['development/','300 per arm; numeric rows, pairs, phases, source comparison and raw ZIP parts'],
 ['full_single_agent/','2,929 per arm; scores, paired and subgroup statistics, diagnostics and baseline raw archive'],
 ['full_v3/','Original execution archive with all 2,931 V3 records and owned audit payloads'],
 ['cost_cache_sensitivity.json','Frozen-rate, fixed-usage cache scenario'],
 ['PACKAGE_MANIFEST.json','Artifact sizes and SHA256 inventory'],
 ['reproduce/','Offline analysis and archive reconstruction helpers']], [2.8,4.3])
h('Metric dictionary',2)
p('question_id is the join key and must be imported as text. Patient IDs define bootstrap clusters. correct is the frozen typed final-answer score. latency_seconds is selected task elapsed time. prompt_tokens includes cached_prompt_tokens; completion_tokens is output usage; total_tokens is prompt plus completion. estimated_cost_usd is the frozen-rate calculation. reported_cost_usd is unavailable provider dollar reporting, not the proxy estimate.')
p('llm_calls and executed-tool counts measure different operations. evidence_count counts recorded evidence entries, while evidence_valid, precision and recall require recorded validation and can be missing. Draft/final transitions apply only where a draft exists. Empty/no-value, limitation and budget-stop flags retain their recorded meanings; none substitutes for typed correctness.')
p('The released benchmark and original V3 execution archive contain 2,931 questions. The report compares the 2,929 question IDs with completed outcomes from both systems. RUN_REGISTRY.json and run_summary.csv retain original execution counts; comparison_summary.csv contains the primary reporting figures. This reporting view preserves the raw records and does not change the benchmark definition.')
page('Offline reproduction')
p('Start from a checkout of results/v3-ful-run. Read METHODS.txt and RUN_REGISTRY.json, then work from the numeric files before restoring the much larger raw material. All commands below are local and require no model credentials.')
h('Verify the package',2)
p('1. Verify PACKAGE_MANIFEST.json and per-archive checksums. Numeric tables work without raw restoration. Blank CSV cells and JSON null mean unavailable. Baseline validation binds matched IDs, source pins, first-clean selection and arithmetic.')
p('2. Follow full_v3/RECONSTRUCTION.txt, full_single_agent/raw/RECONSTRUCTION.txt and development/raw/README.txt. Development ZIP parts reconstruct exactly. Full-arm XZ envelopes preserve parsed canonical rows and decompressed audit text; original whitespace and gzip headers are not recreated. Standard Python JSON and LZMA libraries suffice.')
h('Read one raw question',2)
p('Use the question index to find the patient chunk and line for a question ID. Each decompressed JSONL envelope has canonical and audit_json fields. canonical retains the original parsed result record; JSON-decode audit_json to inspect its saved trace and resources. The index binds row semantics, original audit payload hashes, question ownership and source ordering. This lets a reviewer inspect a discordant pair without reconstructing the whole corpus.')
h('Verify or reconstruct V3',2)
p('From the repository root, run the following with Python 3.10 or newer. Verification streams every chunk and checks hashes and ownership. Reconstruction requires a new output directory outside the archive and sufficient free disk space; it writes the canonical array and fresh gzip containers with the original audit payload bytes.')
for command in [
 'python3 results/v3/full_v3/archive_tools.py verify --archive results/v3/full_v3/data_archive',
 'python3 results/v3/full_v3/archive_tools.py reconstruct --archive results/v3/full_v3/data_archive --output /tmp/v3-restored-results']:
    par=p(command)
    for run in par.runs: run.font.name='Consolas';run.font.size=Pt(9)
p('The baseline archive uses the same envelope schema. Its own reconstruction guide gives the corresponding utility and archive-directory arguments. Development uses ordered ZIP parts; its guide supplies the exact concatenation order and original SHA256.')
h('Rebuild and interpret tables',2)
p('3. Run reproduce/build_analysis.py with Python and python-docx to rebuild this handoff from saved summaries. No provider request is made. Frozen scorer and analysis sources remain pinned in Git; derived extraction sources and hashes accompany the numerical packages.')
p('4. Use comparison_summary.csv and the matched 2,929-question files for full benchmark figures. Use development files for the reused checkpoint. The original V3 execution package supplies archival provenance. Full comparison paired JSON files use Single Agent minus V3; development paired tables and this handoff use V3 minus Single Agent.')
p('Before publication, verify that figure labels preserve these denominators and signs. Recomputing an estimate from a selected subset creates a new analysis; retain its question IDs and assumptions alongside the resulting figure. The archived raw and numerical inputs remain the reference records.')

doc.core_properties.title='V3 consolidated results and analysis'
doc.core_properties.subject='Midterm benchmark checkpoint and reproducible handoff'
doc.core_properties.author='CSE 5106 project team'
doc.core_properties.keywords='V3, FHIR, benchmark, paired comparison, midterm'
doc.save(BASE / 'ANALYSIS.docx')
(BASE / 'ANALYSIS.txt').write_text('V3 CONSOLIDATED RESULTS\nMidterm checkpoint and analysis handoff\nRevised 19 October 2026\n\n'+'\n'.join(text))

comparison_rows=[]
for name,arm in [('single_agent',s),('multi_agent_v3',v)]:
    row={'arm':name,'questions':arm['questions'],'patients':arm['accuracy']['patients'],
         'correct':arm['correct'],'wrong':arm['questions']-arm['correct'],
         'accuracy':arm['accuracy']['estimate'],'accuracy_ci_low':arm['accuracy']['ci_low'],
         'accuracy_ci_high':arm['accuracy']['ci_high']}
    for key in ['latency_seconds','total_tokens','estimated_cost_usd']:
        for stat in ['total','mean','median','p95','stdev']:
            row[key+'_'+stat]=arm['metrics'][key][stat]
    comparison_rows.append(row)
with (BASE/'comparison_summary.csv').open('w',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=list(comparison_rows[0]),lineterminator='\n');writer.writeheader()
    writer.writerows(comparison_rows)

with (BASE/'run_summary.csv').open('w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=['run_id','planned_n','recorded_n','correct','accuracy','selection_policy']);w.writeheader()
    for r in registry['runs']:
        w.writerow(dict(run_id=r['id'],planned_n=r['planned_n'],recorded_n=r['accepted_or_scheduled_n'],correct=r['correct'],accuracy=r['correct']/r['accepted_or_scheduled_n'],selection_policy=r['policy']))
print(json.dumps({'docx':str(BASE/'ANALYSIS.docx'),'pages_planned':11,'text_words':len(' '.join(text).split()),'source_rows':[300,300,2931,2929]}))
